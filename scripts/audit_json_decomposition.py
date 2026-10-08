"""以 2017.xlsx 为唯一原文依据，校核并生成可追溯 JSON 数据集。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from copy import deepcopy
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL = PROJECT_ROOT / "datasets" / "lunwen" / "data" / "2017.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "datasets" / "lunwen" / "verified_json"


@dataclass(slots=True)
class AuditRecord:
    json_file: str
    excel_sheet: str
    excel_row: int
    source_row_id: str
    clause_index: int
    source_text: str
    selected_source: str
    source_sha256: str
    output_sha256: str
    status: str
    quality_score: float
    corrected: bool
    manual_review_required: bool
    reasons: str
    rules_before: str
    rules_after: str
    entities_before: str
    entities_after: str
    relations_before: int
    relations_after: int
    split_content_mismatch: bool


def normalize(text: Any) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]", "", str(text or "")).lower()


def hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def best_coverage(fragment: str, source: str) -> float:
    """片段在原文中的最长连续字符覆盖率，适合检测明显串行错位。"""
    left, right = normalize(fragment), normalize(source)
    if not left or not right:
        return 0.0
    if left in right:
        return 1.0
    match = SequenceMatcher(None, left, right, autojunk=False).find_longest_match()
    return match.size / len(left)


def safe_rule(source_text: str, entity_names: list[str]) -> dict[str, Any]:
    return {
        "content": source_text,
        "type": "待人工复核",
        "subject": "",
        "action": "",
        "condition": "",
        "constraint": "",
        "related_entities": entity_names,
    }


def source_rows(excel_path: Path) -> list[tuple[int, str, str, str]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = []
    for excel_row, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        row_id, clause1, clause2 = (values + (None, None, None))[:3]
        rows.append((excel_row, str(row_id or ""), str(clause1 or "").strip(), str(clause2 or "").strip()))
    return rows


def collect_sources(root: Path) -> dict[str, list[tuple[str, Path, str]]]:
    locations = [
        ("output_jsons/train", root / "output_jsons" / "train"),
        ("output_jsons/test", root / "output_jsons" / "test"),
        ("output_jsons_new/train", root / "output_jsons_new" / "train"),
        ("output_jsons_new/test", root / "output_jsons_new" / "test"),
    ]
    grouped: dict[str, list[tuple[str, Path, str]]] = {}
    for label, directory in locations:
        for path in directory.glob("*.json"):
            raw = path.read_text(encoding="utf-8")
            grouped.setdefault(path.name, []).append((label, path, raw))
    return grouped


def choose_source(candidates: list[tuple[str, Path, str]]) -> tuple[str, Path, str, bool]:
    """优先采用 output_jsons_new，保留不同划分内容不一致标识。"""
    hashes = {hash_text(raw) for _, _, raw in candidates}
    ordered = sorted(candidates, key=lambda item: (not item[0].startswith("output_jsons_new"), item[0]))
    label, path, raw = ordered[0]
    return label, path, raw, len(hashes) > 1


def audit_payload(payload: dict[str, Any], source_text: str) -> tuple[dict[str, Any], str, float, bool, bool, list[str]]:
    """仅自动修正高置信度错误：缺失/坏 JSON 或规则内容明显无关。"""
    reasons: list[str] = []
    corrected = False
    manual = False
    entities = payload.get("entities", [])
    rules = payload.get("rules", [])
    relations = payload.get("relations", [])
    if not isinstance(entities, list) or not isinstance(rules, list) or not isinstance(relations, list):
        reasons.append("schema_fields_not_lists")
        entities, rules, relations = [], [], []
        corrected = True
    rule_contents = [str(item.get("content", "")) for item in rules if isinstance(item, dict)]
    coverages = [best_coverage(item, source_text) for item in rule_contents if normalize(item)]
    if not rule_contents:
        reasons.append("missing_rules")
        corrected = True
    elif any(score < 0.20 for score in coverages):
        reasons.append("rule_content_unrelated_to_excel_source")
        corrected = True
    elif any(score < 0.75 for score in coverages):
        reasons.append("rule_content_requires_manual_review")
        manual = True

    entity_names = [str(item.get("name", "")) for item in entities if isinstance(item, dict) and item.get("name")]
    entity_scores = [best_coverage(name, source_text) for name in entity_names if normalize(name)]
    if entity_names and sum(score >= 0.75 for score in entity_scores) / len(entity_names) < 0.5:
        reasons.append("entity_grounding_requires_manual_review")
        manual = True

    if corrected:
        grounded_entities = [item for item in entities if isinstance(item, dict) and item.get("name") and best_coverage(str(item["name"]), source_text) >= 0.75]
        entity_names = [str(item["name"]) for item in grounded_entities]
        payload = {
            "entities": grounded_entities,
            "rules": [safe_rule(source_text, entity_names)],
            "relations": [],
        }
        status = "corrected_to_source_faithful_rule"
        score = 1.0
    else:
        status = "needs_manual_review" if manual else "passed_source_grounding"
        rule_score = sum(coverages) / len(coverages) if coverages else 0.0
        entity_score = sum(entity_scores) / len(entity_scores) if entity_scores else 1.0
        score = round(0.75 * rule_score + 0.25 * entity_score, 4)
    return payload, status, score, corrected, manual, reasons


def create_missing_payload(source_text: str) -> dict[str, Any]:
    return {"entities": [], "rules": [safe_rule(source_text, [])], "relations": []}


def add_metadata(payload: dict[str, Any], *, excel_row: int, source_row_id: str, clause_index: int, source_text: str, status: str, score: float, corrected: bool, manual: bool, reasons: list[str], selected_source: str) -> dict[str, Any]:
    output = deepcopy(payload)
    output["metadata"] = {
        "source": {
            "excel_file": "2017.xlsx",
            "excel_sheet": "Sheet1",
            "excel_row": excel_row,
            "source_row_id": source_row_id,
            "clause_index": clause_index,
            "source_text": source_text,
        },
        "audit": {
            "status": status,
            "quality_score": score,
            "corrected": corrected,
            "manual_review_required": manual,
            "reasons": reasons,
            "selected_source": selected_source,
            "audited_at": datetime.now(timezone.utc).isoformat(),
            "audit_version": "1.0",
        },
    }
    return output


def write_mapping_xlsx(records: list[AuditRecord], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "修正对照表"
    headers = list(asdict(records[0]).keys()) if records else list(AuditRecord.__annotations__)
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for record in records:
        sheet.append(list(asdict(record).values()))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        width = 18
        if header in {"source_text", "rules_before", "rules_after", "entities_before", "entities_after", "reasons"}:
            width = 48
        elif header in {"source_sha256", "output_sha256"}:
            width = 20
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    summary = workbook.create_sheet("汇总")
    summary.append(["指标", "值"])
    summary.append(["JSON总数", len(records)])
    for status, count in sorted(Counter(record.status for record in records).items()):
        summary.append([f"状态：{status}", count])
    summary.append(["自动修正", sum(record.corrected for record in records)])
    summary.append(["需人工复核", sum(record.manual_review_required for record in records)])
    for cell in summary[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = header_fill
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 18
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_readme(output_dir: Path, summary: dict[str, Any]) -> None:
    text = "# JSON分解校核结果\n\n本目录以 `datasets/lunwen/data/2017.xlsx` 的 Sheet1 为原文依据。文件名 `N_clause1.json` 中的 N 是 Sheet1 第一列的条款对编号；`N_clause1.json` 表示该编号对应行的第一个条款，`N_clause2.json` 表示第二个条款。Excel 的物理行号单独保存在 `metadata.source.excel_row` 中。\n\n"
    text += "## 生成内容\n\n- `json/`：每个 Excel 条款对应一个已校核 JSON；包含 `metadata.source` 和 `metadata.audit`。\n- `修正对照表.xlsx`：完整 Excel 对照表与汇总页。\n- `修正对照表.csv`：同一内容的文本表。\n- `audit_summary.json`：可程序读取的统计。\n\n"
    text += "## 本次统计\n\n```json\n" + json.dumps(summary, ensure_ascii=False, indent=2) + "\n```\n\n"
    text += "自动修正只处理文件缺失、JSON结构损坏或规则文本与原文明显无关的高置信度错误；语义近似、实体别名或轻微改写被标为 `needs_manual_review`，不自动删除。\n"
    (output_dir / "README.md").write_text(text, encoding="utf-8")


def run(excel_path: Path, dataset_root: Path, output_dir: Path, clean: bool = False) -> dict[str, Any]:
    rows = source_rows(excel_path)
    source_map = collect_sources(dataset_root)
    json_dir = output_dir / "json"
    json_dir.mkdir(parents=True, exist_ok=True)
    expected_names = {f"{source_row_id}_clause{clause_index}.json" for _, source_row_id, clause1, clause2 in rows for clause_index, source_text in ((1, clause1), (2, clause2)) if source_text}
    if clean:
        for path in json_dir.glob("*_clause*.json"):
            if path.name not in expected_names:
                path.unlink()
    records: list[AuditRecord] = []
    for excel_row, source_row_id, clause1, clause2 in rows:
        for clause_index, source_text in ((1, clause1), (2, clause2)):
            name = f"{source_row_id}_clause{clause_index}.json"
            if not source_text:
                records.append(AuditRecord(
                    json_file="", excel_sheet="Sheet1", excel_row=excel_row, source_row_id=source_row_id,
                    clause_index=clause_index, source_text="", selected_source="", source_sha256="", output_sha256="",
                    status="skipped_blank_source_text", quality_score=1.0, corrected=False, manual_review_required=False,
                    reasons="blank_source_text", rules_before="", rules_after="", entities_before="", entities_after="",
                    relations_before=0, relations_after=0, split_content_mismatch=False,
                ))
                continue
            candidates = source_map.get(name, [])
            selected_source = ""
            split_mismatch = False
            original_raw = ""
            reasons: list[str] = []
            if not candidates:
                payload = create_missing_payload(source_text)
                status, score, corrected, manual = "created_missing_decomposition", 1.0, True, True
                reasons.append("missing_in_all_original_json_splits")
                rules_before, entities_before, relations_before = "", "", 0
            else:
                selected_source, _, original_raw, split_mismatch = choose_source(candidates)
                try:
                    original = json.loads(original_raw)
                    rules_before = json.dumps(original.get("rules", []), ensure_ascii=False)
                    entities_before = json.dumps(original.get("entities", []), ensure_ascii=False)
                    relations_before = len(original.get("relations", [])) if isinstance(original.get("relations", []), list) else 0
                    payload, status, score, corrected, manual, reasons = audit_payload(original, source_text)
                except (json.JSONDecodeError, AttributeError):
                    payload = create_missing_payload(source_text)
                    status, score, corrected, manual = "repaired_invalid_json", 1.0, True, True
                    reasons.append("invalid_json")
                    rules_before, entities_before, relations_before = original_raw, "", 0
            if split_mismatch:
                reasons.append("split_content_mismatch")
                manual = True
            output = add_metadata(payload, excel_row=excel_row, source_row_id=source_row_id, clause_index=clause_index, source_text=source_text, status=status, score=score, corrected=corrected, manual=manual, reasons=reasons, selected_source=selected_source)
            serialized = json.dumps(output, ensure_ascii=False, indent=2)
            (json_dir / name).write_text(serialized, encoding="utf-8")
            records.append(AuditRecord(
                json_file=f"json/{name}", excel_sheet="Sheet1", excel_row=excel_row, source_row_id=source_row_id,
                clause_index=clause_index, source_text=source_text, selected_source=selected_source,
                source_sha256=hash_text(original_raw) if original_raw else "", output_sha256=hash_text(serialized), status=status,
                quality_score=score, corrected=corrected, manual_review_required=manual, reasons="; ".join(reasons),
                rules_before=rules_before, rules_after=json.dumps(output["rules"], ensure_ascii=False),
                entities_before=entities_before, entities_after=json.dumps(output["entities"], ensure_ascii=False),
                relations_before=relations_before, relations_after=len(output["relations"]), split_content_mismatch=split_mismatch,
            ))
    summary = {
        "excel_file": str(excel_path), "excel_rows": len(rows), "source_clause_cells": len(records),
        "expected_json_files": len(expected_names), "generated_json_files": sum(bool(record.json_file) for record in records),
        "original_json_files_unique": len(source_map), "status_counts": dict(Counter(record.status for record in records)),
        "corrected_count": sum(record.corrected for record in records), "manual_review_count": sum(record.manual_review_required for record in records),
        "split_content_mismatch_count": sum(record.split_content_mismatch for record in records),
        "generated_at": datetime.now(timezone.utc).isoformat(), "audit_version": "1.0",
    }
    with (output_dir / "修正对照表.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(asdict(records[0]).keys()))
        writer.writeheader()
        writer.writerows(asdict(record) for record in records)
    write_mapping_xlsx(records, output_dir / "修正对照表.xlsx")
    (output_dir / "audit_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_readme(output_dir, summary)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="校核2017.xlsx与JSON分解结果")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--dataset-root", type=Path, default=PROJECT_ROOT / "datasets" / "lunwen")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--clean", action="store_true", help="删除输出目录中不再对应Excel条款对编号的旧JSON")
    args = parser.parse_args()
    print(json.dumps(run(args.excel, args.dataset_root, args.output, args.clean), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
