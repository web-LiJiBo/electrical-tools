"""Read supplemental Excel standards without changing the original knowledge base.

Each nonempty selected worksheet: row 1 = name, row 2 = identifier,
row 3 onward = one subclause per nonempty row, as specified by the user.
No formulas/macros are executed and no workbook is saved.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime, time, timezone
import hashlib
import io
import json
from pathlib import Path
import uuid


# Repository-contained default; original external working data is preserved separately.
DEFAULT_ROOT = Path(__file__).resolve().parents[3] / "datasets" / "supplemental"
SUPPORTED = {".xlsx", ".xlsm"}


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def _header(sheet, row: int) -> str:
    values = []
    for cell in sheet[row]:
        if cell.value is None or not str(cell.value).strip():
            continue
        if cell.data_type in {"f", "e"} or not isinstance(cell.value, str):
            raise ValueError(f"第{row}行须为文本，不能使用公式、错误值或纯数字")
        values.append(cell.value.strip())
    unique = list(dict.fromkeys(values))
    if len(unique) != 1:
        raise ValueError(f"第{row}行须有一个明确的标准{'名称' if row == 1 else '编号'}，当前有{len(unique)}个候选值")
    return unique[0]


def read_excel(path: Path, *, source_key: str | None = None,
               sheets: list[str] | None = None) -> dict:
    """Return raw rows and provenance. Invalid sheets are reported individually."""
    import openpyxl

    path = Path(path).resolve()
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError("仅支持.xlsx/.xlsm；旧版.xls请另存为.xlsx后读取")
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    source_key = source_key or path.name
    document_id = "supp_excel_" + hashlib.sha256(
        (source_key + "\0" + digest).encode("utf-8")
    ).hexdigest()
    book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cached = None
    result = {
        "schema_version": 1, "document_id": document_id,
        "dataset": "supplemental_excel", "stage": "raw_import",
        "semantic_status": "not_parsed", "approved_for_control": False,
        "source_file": str(path), "source_key": source_key,
        "source_sha256": digest, "source_size": len(raw),
        "sheets": [], "errors": [], "skipped_sheets": [],
    }
    try:
        cached = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
        selected = sheets if sheets is not None else book.sheetnames
        for name in selected:
            if name not in book.sheetnames:
                result["errors"].append({"sheet": name, "message": "工作表不存在"})
                continue
            sheet = book[name]
            if sheet.max_row > 100000 or sheet.max_column > 256:
                result["errors"].append({"sheet": name, "message": "工作表范围过大（上限100000行、256列），请清理多余格式后重试"})
                continue
            if not any(c.value is not None for row in sheet.iter_rows() for c in row):
                result["skipped_sheets"].append({"sheet": name, "reason": "空工作表"})
                continue
            try:
                standard_name, standard_id = _header(sheet, 1), _header(sheet, 2)
                rows, warnings = [], []
                for row in sheet.iter_rows(min_row=3):
                    cells = []
                    for cell in row:
                        if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                            continue
                        value = cell.value
                        item = {"coordinate": cell.coordinate,
                                "value": _json_value(value),
                                "value_type": type(value).__name__,
                                "excel_type": cell.data_type,
                                "number_format": cell.number_format}
                        if cell.data_type == "f":
                            formula = value if isinstance(value, str) else getattr(value, "text", str(value))
                            item["value"] = formula
                            item["cached_value"] = _json_value(cached[name][cell.coordinate].value)
                            item["cache_verified"] = False
                            warnings.append(f"{cell.coordinate}为公式，缓存未校验，不能作为已验证条款值")
                        if cell.data_type == "e":
                            warnings.append(f"{cell.coordinate}含Excel错误值：{value}")
                        cells.append(item)
                    if cells:
                        source_row = row[0].row
                        rows.append({
                            "row_id": f"{document_id}_{hashlib.sha256(name.encode()).hexdigest()[:16]}_r{source_row}",
                            "source_row": source_row, "unit_kind": "subclause",
                            "raw_text": "\t".join(str(c["value"]) for c in cells),
                            "cells": cells,
                        })
                if not rows:
                    raise ValueError("第3行起没有条款正文")
                result["sheets"].append({
                    "sheet_name": name, "sheet_state": sheet.sheet_state,
                    "standard_name": standard_name, "standard_id": standard_id,
                    "header_rows": [1, 2], "merged_ranges": [str(r) for r in sheet.merged_cells.ranges],
                    "rows": rows, "warnings": warnings,
                })
            except ValueError as exc:
                result["errors"].append({"sheet": name, "message": str(exc)})
        if not result["sheets"] and not result["errors"]:
            result["errors"].append({"sheet": None, "message": "没有可读取的工作表"})
        return result
    finally:
        book.close()
        if cached is not None:
            cached.close()


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def build_subclause_index(report_path: Path, destination: Path) -> dict:
    """Index the latest imported snapshot: one JSONL record per source row.

    This is citable source evidence, not parsed executable rules. Never scan all
    historic documents, guess the official subclause number, or invent actions.
    """
    report_path, destination = Path(report_path).resolve(), Path(destination).resolve()
    report = json.loads(report_path.read_text(encoding="utf-8"))
    input_dir = Path(report["input_dir"]).resolve()
    if destination.is_relative_to(input_dir) or input_dir.is_relative_to(destination.parent):
        raise ValueError("子条款索引不能写入或覆盖原始Excel目录")
    if destination == report_path or destination.suffix.lower() != ".jsonl":
        raise ValueError("请使用独立的.jsonl索引文件")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + "." + uuid.uuid4().hex + ".tmp")
    count = 0
    identifiers = set()
    try:
        with temporary.open("w", encoding="utf-8") as stream:
            for document in report["documents"]:
                payload = json.loads(Path(document["path"]).read_text(encoding="utf-8"))
                for sheet in payload["sheets"]:
                    for row in sheet["rows"]:
                        clause_id = row["row_id"]
                        if clause_id in identifiers:
                            raise ValueError(f"重复子条款标识：{clause_id}")
                        identifiers.add(clause_id)
                        unresolved = any(c.get("excel_type") in {"f", "e"} for c in row["cells"])
                        if unresolved:
                            raise ValueError(f"{sheet['standard_id']} {sheet['sheet_name']} 第{row['source_row']}行含未校验公式或错误值，请先核实后导出")
                        unit = {
                            "standard_name": sheet["standard_name"],
                            "standard_id": sheet["standard_id"],
                            # Keep the source's clause number verbatim in the text.
                            # Do not generate a new number or split/renumber a row.
                            "text": row["raw_text"],
                        }
                        stream.write(json.dumps(unit, ensure_ascii=False) + "\n")
                        count += 1
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    summary = {"source_run_id": report["run_id"], "document_count": len(report["documents"]),
               "subclause_count": count,
               "failed_source_count": len({e["file"] for e in report["errors"]}),
               "index_path": str(destination), "row_convention": "one_row_one_clause",
               "connected_to_assessment_pipeline": False}
    _write_json(destination.with_suffix(".summary.json"), summary)
    return summary


def import_directory(input_dir: Path, output_dir: Path, report_dir: Path,
                     *, sheets: list[str] | None = None) -> dict:
    input_dir, output_dir, report_dir = (Path(p).resolve() for p in (input_dir, output_dir, report_dir))
    for destination in (output_dir, report_dir):
        if destination == input_dir or destination.is_relative_to(input_dir) or input_dir.is_relative_to(destination):
            raise ValueError("输入目录与结果/记录目录不能相互包含")
    if output_dir == report_dir or output_dir.is_relative_to(report_dir) or report_dir.is_relative_to(output_dir):
        raise ValueError("读取结果与导入记录必须放在独立目录")
    input_dir.mkdir(parents=True, exist_ok=True)
    report = {"schema_version": 1, "run_id": uuid.uuid4().hex,
              "imported_at": datetime.now(timezone.utc).isoformat(),
              "input_dir": str(input_dir), "documents": [], "errors": [],
              "row_count": 0, "warning_count": 0}
    files = sorted(p for p in input_dir.rglob("*") if p.is_file()
                   and p.suffix.lower() in SUPPORTED | {".xls"} and not p.name.startswith("~$"))
    for path in files:
        try:
            payload = read_excel(path, source_key=path.relative_to(input_dir).as_posix(), sheets=sheets)
            report["errors"].extend({"file": str(path), **e} for e in payload["errors"])
            if payload["sheets"]:
                target = output_dir / (payload["document_id"] + ".json")
                _write_json(target, payload)
                report["documents"].append({"document_id": payload["document_id"], "path": str(target),
                                            "source_file": str(path), "sha256": payload["source_sha256"]})
                report["row_count"] += sum(len(s["rows"]) for s in payload["sheets"])
                report["warning_count"] += sum(len(s["warnings"]) for s in payload["sheets"])
        except Exception as exc:
            # A broken file must not prevent other independent files from being read.
            report["errors"].append({"file": str(path), "message": f"{type(exc).__name__}: {exc}"})
    report["file_count"] = len(files)
    report["status"] = "failed_or_partial" if report["errors"] else ("success" if files else "no_input")
    _write_json(report_dir / (report["run_id"] + ".json"), report)
    # Current snapshot only; old content-addressed documents remain as history.
    _write_json(report_dir / "latest.json", report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="独立读取补充标准Excel，不修改原标准库")
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_ROOT / "原始Excel")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_ROOT / "读取结果")
    parser.add_argument("--report-dir", type=Path, default=DEFAULT_ROOT / "导入记录")
    parser.add_argument("--sheet", action="append", help="仅读取指定工作表，可重复使用")
    parser.add_argument("--build-subclauses", action="store_true", help="将已读取的最新清单转为逐行子条款，不重复读取Excel")
    args = parser.parse_args()
    try:
        if args.build_subclauses:
            summary = build_subclause_index(args.report_dir / "latest.json",
                                           args.output_dir.parent / "子条款" / "子条款索引.jsonl")
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            return 0
        report = import_directory(args.input_dir, args.output_dir, args.report_dir, sheets=args.sheet)
    except (ValueError, OSError, ImportError) as exc:
        print(f"读取失败：{exc}")
        return 1
    print(f"读取文件：{report['file_count']}，成功文档：{len(report['documents'])}，正文行：{report['row_count']}，警告：{report['warning_count']}，错误：{len(report['errors'])}")
    if report["status"] == "no_input":
        print(f"尚无Excel文件，请放入：{args.input_dir}")
    for error in report["errors"]:
        print(json.dumps(error, ensure_ascii=False))
    print(f"导入记录：{args.report_dir / 'latest.json'}")
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
